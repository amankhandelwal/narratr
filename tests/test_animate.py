"""On-click reveals as timing XML, in the shape Google Slides was seen to import."""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from narratr.deck import animate


def slide_with(n: int):
	prs = Presentation()
	slide = prs.slides.add_slide(prs.slide_layouts[6])
	ids = [
		slide.shapes.add_textbox(Inches(i), Inches(1), Inches(1), Inches(1)).shape_id
		for i in range(n)
	]
	return prs, slide, ids


def test_groups_round_trip():
	_, slide, ids = slide_with(4)
	groups = [[ids[0]], [ids[1], ids[2]], [ids[3]]]
	animate.apply(slide, groups)
	assert animate.click_groups(slide._element) == groups


def test_first_of_a_group_is_on_click_the_rest_with_previous():
	xml = animate.timing_xml([[5, 6, 7]])
	assert xml.count('nodeType="clickEffect"') == 1
	assert xml.count('nodeType="withEffect"') == 2


def test_every_effect_is_a_fade():
	xml = animate.timing_xml([[5], [6]])
	assert xml.count('presetID="10" presetClass="entr"') == 2
	assert xml.count('filter="fade"') == 2


def test_no_groups_writes_no_timing():
	_, slide, _ = slide_with(1)
	animate.apply(slide, [])
	assert animate.click_groups(slide._element) == []
	assert slide._element.find(qn("p:timing")) is None


def test_timing_ids_are_unique():
	xml = animate.timing_xml([[1, 2], [3], [4, 5, 6]])
	ids = re.findall(r'<p:cTn id="(\d+)"', xml)
	assert len(ids) == len(set(ids))


def test_saved_file_reopens(tmp_path):
	prs, slide, ids = slide_with(2)
	animate.apply(slide, [[ids[0]], [ids[1]]])
	prs.save(tmp_path / "x.pptx")
	again = Presentation(tmp_path / "x.pptx")
	assert animate.click_groups(again.slides[0]._element) == [[ids[0]], [ids[1]]]
