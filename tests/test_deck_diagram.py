"""A flowchart becomes native shapes: one per node, one glued connector per edge."""

from __future__ import annotations

import json
import re

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from narratr.deck import animate
from narratr.deck.build import build
from narratr.paths import ROOT
from narratr.spec import is_flowchart

SPECS = [
	ROOT / "examples" / "brief.scenes.json",
	ROOT / "skills" / "narratr" / "example.scenes.json",
]
NODE = re.compile(r"^\s*([A-Za-z0-9_]+)\s*(?:@\{|\[|\(|\{)", re.M)
EDGE = re.compile(r"^\s*([A-Za-z0-9_]+)\s*-->(?:\|[^|]*\|)?\s*([A-Za-z0-9_]+)", re.M)


def diagram_slides():
	for path in SPECS:
		spec = json.loads(path.read_text())
		for index, scene in enumerate(spec["scenes"]):
			if scene["type"] == "diagram" and is_flowchart(scene["mermaid"]):
				yield pytest.param(path, index, id=f"{path.stem}:{scene['id']}")


CASES = list(diagram_slides())


@pytest.fixture(scope="module")
def decks(tmp_path_factory):
	out = {}
	for path in SPECS:
		spec = json.loads(path.read_text())
		deck, _ = build(spec, path.parent, tmp_path_factory.mktemp(path.stem))
		out[path] = (spec, Presentation(str(deck)))
	return out


def test_there_are_diagrams_to_check():
	assert len(CASES) >= 2


@pytest.mark.parametrize(("path", "index"), CASES)
def test_every_node_is_exactly_one_shape(decks, path, index):
	spec, prs = decks[path]
	scene, slide = spec["scenes"][index], prs.slides[index]
	ids = set(NODE.findall(scene["mermaid"]))
	names = [s.name for s in slide.shapes if s.name.startswith("node:")]
	assert sorted(names) == sorted(f"node:{i}" for i in ids)


@pytest.mark.parametrize(("path", "index"), CASES)
def test_every_edge_is_a_connector_glued_to_its_nodes(decks, path, index):
	spec, prs = decks[path]
	scene, slide = spec["scenes"][index], prs.slides[index]
	by_name = {s.name: s for s in slide.shapes}
	edges = EDGE.findall(scene["mermaid"])
	assert edges
	for src, dst in edges:
		edge = by_name[f"edge:{src}->{dst}"]
		start = edge._element.find(".//{*}stCxn").get("id")
		end = edge._element.find(".//{*}endCxn").get("id")
		assert int(start) == by_name[f"node:{src}"].shape_id
		assert int(end) == by_name[f"node:{dst}"].shape_id


@pytest.mark.parametrize(("path", "index"), CASES)
def test_a_reveal_step_is_one_click_and_brings_its_edges(decks, path, index):
	spec, prs = decks[path]
	scene, slide = spec["scenes"][index], prs.slides[index]
	by_id = {s.shape_id: s.name for s in slide.shapes}
	groups = animate.click_groups(slide._element)
	assert len(groups) == len(scene["revealSteps"])
	beat = {n: i for i, step in enumerate(scene["revealSteps"]) for n in step}
	for i, group in enumerate(groups):
		names = {by_id[s] for s in group}
		assert {f"node:{n}" for n in scene["revealSteps"][i]} <= names
		for src, dst in EDGE.findall(scene["mermaid"]):
			if max(beat.get(src, -1), beat.get(dst, -1)) == i:
				assert f"edge:{src}->{dst}" in names


@pytest.mark.parametrize(("path", "index"), CASES)
def test_node_labels_are_editable_text(decks, path, index):
	spec, prs = decks[path]
	slide = prs.slides[index]
	texts = {s.text_frame.text for s in slide.shapes if s.has_text_frame}
	labels = re.findall(r'label:\s*"([^"]+)"', spec["scenes"][index]["mermaid"])
	assert labels
	assert set(labels) <= texts


def test_underscored_node_ids_resolve_their_edges(tmp_path):
	"""Mermaid joins edge ids with underscores, which node ids may contain too."""
	spec = json.loads(SPECS[0].read_text())
	scene = next(s for s in spec["scenes"] if s["type"] == "diagram")
	scene["mermaid"] = (
		'flowchart LR\n    step_one["One"]\n    step_two["Two"]\n    step_one --> step_two'
	)
	scene["revealSteps"] = [["step_one"], ["step_two"]]
	scene["revealCues"] = scene["revealCues"][:2]
	deck, _ = build(spec, SPECS[0].parent, tmp_path)
	slide = Presentation(str(deck)).slides[spec["scenes"].index(scene)]
	assert "edge:step_one->step_two" in {s.name for s in slide.shapes}


def test_a_sequence_diagram_is_one_static_picture(tmp_path):
	spec = json.loads(SPECS[0].read_text())
	scene = next(s for s in spec["scenes"] if s["type"] == "diagram")
	scene["mermaid"] = "sequenceDiagram\n    A->>B: hello"
	deck, _ = build(spec, SPECS[0].parent, tmp_path)
	slide = Presentation(str(deck)).slides[spec["scenes"].index(scene)]
	assert animate.click_groups(slide._element) == []
	assert any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
