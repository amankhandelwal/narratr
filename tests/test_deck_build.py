"""Properties of a built deck, over every slide of every bundled example."""

from __future__ import annotations

import json

import pytest
from pptx import Presentation

from narratr.deck import animate, script
from narratr.deck.build import build
from narratr.deck.layout import LAYOUTS
from narratr.paths import ROOT

SPECS = {
	"brief": ROOT / "examples" / "brief.scenes.json",
	"code": ROOT / "examples" / "code.scenes.json",
	"skill": ROOT / "skills" / "narratr" / "example.scenes.json",
}


@pytest.fixture(scope="module", params=sorted(SPECS))
def built(request, tmp_path_factory):
	path = SPECS[request.param]
	spec = json.loads(path.read_text())
	deck, doc = build(spec, path.parent, tmp_path_factory.mktemp(request.param))
	return spec, Presentation(str(deck)), doc.read_text()


def shape_ids(slide) -> set[int]:
	return {int(e.get("id")) for e in slide._element.iter() if e.tag.endswith("}cNvPr")}


def test_one_slide_per_scene(built):
	spec, prs, _ = built
	assert len(prs.slides) == len(spec["scenes"])


def test_clicks_agree_three_ways(built):
	"""Timing groups, spec reveals and notes markers are the same count."""
	spec, prs, _ = built
	for scene, slide in zip(spec["scenes"], prs.slides, strict=True):
		groups = animate.click_groups(slide._element)
		notes = slide.notes_slide.notes_text_frame.text
		assert len(groups) == len(script.reveals(scene)), scene["id"]
		assert notes.count(script.MARKER) == len(groups), scene["id"]


def test_every_animation_targets_a_shape_on_its_slide(built):
	_, prs, _ = built
	for slide in prs.slides:
		ids = shape_ids(slide)
		for group in animate.click_groups(slide._element):
			assert set(group) <= ids


def test_notes_are_the_narration_with_markers(built):
	spec, prs, _ = built
	for scene, slide in zip(spec["scenes"], prs.slides, strict=True):
		notes = slide.notes_slide.notes_text_frame.text
		assert notes.replace(script.MARKER + " ", "") == scene["narration"]


def test_script_md_matches_the_notes(built):
	_, prs, doc = built
	for slide in prs.slides:
		assert slide.notes_slide.notes_text_frame.text in doc


def test_headings_are_title_placeholders(built):
	spec, prs, _ = built
	for scene, slide in zip(spec["scenes"], prs.slides, strict=True):
		if scene.get("heading"):
			assert slide.shapes.title is not None, scene["id"]
			assert slide.shapes.title.text_frame.text == scene["heading"]
		else:
			assert slide.shapes.title is None, scene["id"]


def test_no_empty_placeholders_are_left(built):
	"""An empty placeholder shows 'Click to add text' to whoever edits the deck."""
	_, prs, _ = built
	for slide in prs.slides:
		for placeholder in slide.placeholders:
			assert placeholder.text_frame.text


def test_no_body_text_below_twenty_points(built):
	spec, prs, _ = built
	for scene, slide in zip(spec["scenes"], prs.slides, strict=True):
		if scene["type"] in {"code", "diagram"}:
			continue
		for e in slide._element.iter():
			if e.tag.endswith("}rPr") and e.get("sz"):
				assert int(e.get("sz")) >= 2000, scene["id"]


def test_cards_and_steps_fade_in_whole(built):
	"""A card or step is one group shape, so it cannot appear half-drawn."""
	from pptx.enum.shapes import MSO_SHAPE_TYPE

	spec, prs, _ = built
	for scene, slide in zip(spec["scenes"], prs.slides, strict=True):
		if scene["type"] not in {"flow", "cards"}:
			continue
		by_id = {s.shape_id: s for s in slide.shapes}
		for group in animate.click_groups(slide._element):
			assert by_id[group[0]].shape_type == MSO_SHAPE_TYPE.GROUP, scene["id"]


def test_the_bundled_examples_exercise_flow_and_cards():
	"""Otherwise the test above passes by checking nothing."""
	kinds = {s["type"] for path in SPECS.values() for s in json.loads(path.read_text())["scenes"]}
	assert {"flow", "cards"} <= kinds & set(LAYOUTS)


def test_a_short_listing_gets_a_short_panel(tmp_path):
	"""A one-line command in a full-height panel reads as a broken slide."""
	from pptx.util import Emu

	from narratr.deck.draw import BODY_H

	path = SPECS["skill"]
	spec = json.loads(path.read_text())
	scene = next(s for s in spec["scenes"] if s["type"] == "code")
	assert scene["code"].count("\n") == 0
	deck, _ = build(spec, path.parent, tmp_path)
	slide = Presentation(str(deck)).slides[spec["scenes"].index(scene)]
	panel = next(s for s in slide.shapes if not s.has_text_frame or not s.text_frame.text)
	assert Emu(panel.height).inches < BODY_H / 2
