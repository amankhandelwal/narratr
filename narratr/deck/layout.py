"""One function per scene shape.

Each draws its slide below the heading and returns the click groups: lists of
shape ids, one list per reveal, in reveal order. The builder checks that count
against the script's markers, so a layout cannot quietly drop a reveal.
"""

from __future__ import annotations

import json
import subprocess
from collections.abc import Callable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from narratr import images
from narratr.deck import glyphs
from narratr.deck.draw import (
	BODY_H,
	BODY_TOP,
	BODY_W,
	MARGIN,
	arrow,
	fit,
	panel,
	picture,
	style,
	text,
	textbox,
)
from narratr.deck.theme import Theme
from narratr.errors import PipelineError
from narratr.render import REMOTION

# Slides rejects very large images on import.
MAX_IMAGE_BYTES = 2 * 1024 * 1024


class LayoutError(PipelineError):
	"""A slide could not be drawn."""


@dataclass(frozen=True)
class Ctx:
	theme: Theme
	base: Path  # image paths are relative to this: the spec's directory
	tmp: Path  # scratch for this build only


Layout = Callable[[Any, dict[str, Any], Ctx], list[list[int]]]


def image_bytes(raw: str, base: Path) -> tuple[bytes, int, int]:
	"""The image as it goes into the deck, plus its pixel size.

	Over 2 MB it is downscaled to 2560 px on the long side, and if that is not
	enough and it has no transparency, re-encoded as JPEG. GIFs pass through:
	resizing one drops every frame but the first. The source is never touched.
	"""
	path = images.resolve(raw, base)
	data = path.read_bytes()
	with Image.open(path) as img:
		width, height = img.size
		if len(data) <= MAX_IMAGE_BYTES or img.format == "GIF":
			return data, width, height
		kind = img.format
		img.thumbnail((2560, 2560))
		out = BytesIO()
		img.save(out, format=kind)
		if out.tell() > MAX_IMAGE_BYTES and img.mode in {"RGB", "L"}:
			out = BytesIO()
			img.convert("RGB").save(out, format="JPEG", quality=85)
		return out.getvalue(), img.size[0], img.size[1]


def prose(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	theme = ctx.theme
	bullets = scene.get("bullets") or []
	width = BODY_W
	if scene.get("image"):
		width = BODY_W * 0.55 if bullets else 0
		data, iw, ih = image_bytes(scene["image"], ctx.base)
		left = MARGIN + (width + 0.4 if bullets else 0)
		picture(slide.shapes, data, *fit(iw, ih, left, BODY_TOP, MARGIN + BODY_W - left, BODY_H))
	if not bullets:
		return []
	row = min(0.95, BODY_H / len(bullets))
	top = BODY_TOP + (BODY_H - row * len(bullets)) / 2
	groups: list[list[int]] = []
	for i, bullet in enumerate(bullets):
		struck = bullet.get("state") == "struck"
		y = top + i * row
		group = slide.shapes.add_group_shape()
		indent = 0.0
		if bullet.get("icon"):
			colour = theme.muted if struck else theme.accent(i)
			icon = glyphs.png(bullet["icon"], colour)
			picture(group.shapes, icon, MARGIN, y + (row - 0.5) / 2, 0.5, 0.5)
			indent = 0.75
		text(
			group.shapes,
			MARGIN + indent,
			y,
			width - indent,
			row,
			bullet["text"],
			size=26,
			colour=theme.muted if struck else theme.text,
			font=theme.body_font,
			strike=struck,
		)
		groups.append([group.shape_id])
	return groups


def tokens(code_: str, lang: str, theme: Theme) -> list[list[list[str]]]:
	"""Shiki's colours, one list of [content, colour] per line."""
	payload = json.dumps(
		{"code": code_, "lang": lang, "theme": "github-light" if theme.is_light else "night-owl"}
	)
	try:
		result = subprocess.run(
			["node", "scripts/tokens.mjs"],
			cwd=REMOTION,
			input=payload,
			capture_output=True,
			text=True,
			timeout=120,
			check=False,
		)
	except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
		raise LayoutError(f"code highlighting failed: {exc}. Run 'narratr doctor'")
	if result.returncode != 0:
		tail = (result.stderr or result.stdout).strip().splitlines()[-4:]
		raise LayoutError("code highlighting failed:\n  " + "\n  ".join(tail))
	return json.loads(result.stdout)


def code(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	theme = ctx.theme
	lines = tokens(scene["code"], scene.get("lang") or "text", theme)
	longest = max((sum(len(t[0]) for t in line) for line in lines), default=1) or 1
	pad = 0.3
	# Monospace glyphs are ~0.6 em wide and lines ~1.25 em tall. Code is the
	# one place text may go below 20 pt: a long listing has to fit somehow.
	size = max(
		12.0,
		min(
			24.0,
			(BODY_H - 2 * pad) * 72 / (max(len(lines), 1) * 1.25),
			(BODY_W - 2 * pad) * 72 / (longest * 0.62),
		),
	)
	panel(slide.shapes, MARGIN, BODY_TOP, BODY_W, BODY_H, fill=theme.panel)
	box = textbox(
		slide.shapes,
		MARGIN + pad,
		BODY_TOP + pad,
		BODY_W - 2 * pad,
		BODY_H - 2 * pad,
		MSO_ANCHOR.TOP,
	)
	frame = box.text_frame
	for n, line in enumerate(lines):
		para = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
		for content, colour in line or [[" ", theme.text]]:
			run = para.add_run()
			run.text = content
			style(run, size=size, colour=colour or theme.text, font=theme.mono_font)
	return []


def image(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	data, iw, ih = image_bytes(scene["image"], ctx.base)
	caption = scene.get("caption")
	height = BODY_H - (0.7 if caption else 0)
	picture(slide.shapes, data, *fit(iw, ih, MARGIN, BODY_TOP, BODY_W, height))
	if caption:
		text(
			slide.shapes,
			MARGIN,
			BODY_TOP + height + 0.1,
			BODY_W,
			0.6,
			caption,
			size=20,
			colour=ctx.theme.muted,
			font=ctx.theme.body_font,
			align=PP_ALIGN.CENTER,
		)
	return []


def flow(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	"""Tiles left to right; the arrow into a step arrives with the step."""
	theme = ctx.theme
	steps = scene["steps"]
	n = len(steps)
	gap = 0.7
	tile_w = min(2.4, (BODY_W - gap * (n - 1)) / n)
	tile_h = 2.4
	total = tile_w * n + gap * (n - 1)
	left = MARGIN + (BODY_W - total) / 2
	footer_h = 0.9 if scene.get("footer") else 0
	top = BODY_TOP + (BODY_H - footer_h - tile_h) / 2
	groups: list[list[int]] = []
	for i, step in enumerate(steps):
		x = left + i * (tile_w + gap)
		group = slide.shapes.add_group_shape()
		panel(group.shapes, x, top, tile_w, tile_h, fill=theme.panel, border=theme.accent(i))
		icon = glyphs.png(step["icon"], theme.accent(i))
		picture(group.shapes, icon, x + (tile_w - 0.9) / 2, top + 0.35, 0.9, 0.9)
		text(
			group.shapes,
			x + 0.1,
			top + 1.35,
			tile_w - 0.2,
			tile_h - 1.45,
			step["label"],
			size=24,
			colour=theme.text,
			font=theme.body_font,
			bold=True,
			align=PP_ALIGN.CENTER,
			anchor=MSO_ANCHOR.TOP,
		)
		ids = [group.shape_id]
		if i:
			mid = top + tile_h / 2
			line = arrow(slide.shapes, x - gap + 0.1, mid, x - 0.1, mid, colour=theme.line)
			ids.append(line.shape_id)
		groups.append(ids)
	footer = scene.get("footer")
	if footer:
		y = BODY_TOP + BODY_H - footer_h
		panel(slide.shapes, left, y, total, footer_h - 0.1, fill=theme.panel)
		icon = glyphs.png(footer["icon"], theme.muted)
		picture(slide.shapes, icon, left + 0.25, y + (footer_h - 0.1 - 0.5) / 2, 0.5, 0.5)
		text(
			slide.shapes,
			left + 0.95,
			y,
			total - 1.2,
			footer_h - 0.1,
			footer["label"],
			size=22,
			colour=theme.muted,
			font=theme.body_font,
		)
	return groups


def cards(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	"""Columns side by side, each fading in whole."""
	theme = ctx.theme
	items = scene["cards"]
	n = len(items)
	gap = 0.4
	w = (BODY_W - gap * (n - 1)) / n
	groups: list[list[int]] = []
	for i, card in enumerate(items):
		x, y = MARGIN + i * (w + gap), BODY_TOP
		group = slide.shapes.add_group_shape()
		panel(group.shapes, x, y, w, BODY_H, fill=theme.panel)
		panel(group.shapes, x, y, w, 0.08, fill=theme.accent(i), shape=MSO_SHAPE.RECTANGLE)
		cursor = y + 0.4
		if card.get("icon"):
			picture(
				group.shapes, glyphs.png(card["icon"], theme.accent(i)), x + 0.35, cursor, 0.7, 0.7
			)
			cursor += 0.95
		text(
			group.shapes,
			x + 0.35,
			cursor,
			w - 0.7,
			1.1,
			card["title"],
			size=26,
			colour=theme.text,
			font=theme.heading_font,
			bold=True,
			anchor=MSO_ANCHOR.TOP,
		)
		cursor += 1.15
		if card.get("detail"):
			text(
				group.shapes,
				x + 0.35,
				cursor,
				w - 0.7,
				1.0,
				card["detail"],
				size=20,
				colour=theme.muted,
				font=theme.body_font,
				anchor=MSO_ANCHOR.TOP,
			)
			cursor += 1.05
		if card.get("image"):
			data, iw, ih = image_bytes(card["image"], ctx.base)
			room = y + BODY_H - 0.25 - cursor
			picture(group.shapes, data, *fit(iw, ih, x + 0.25, cursor, w - 0.5, room))
		groups.append([group.shape_id])
	return groups


LAYOUTS: dict[str, Layout] = {
	"prose": prose,
	"flow": flow,
	"cards": cards,
	"code": code,
	"image": image,
}
