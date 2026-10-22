"""Mermaid flowcharts as native, editable shapes.

Mermaid still does the layout: mmdc renders the SVG exactly as the video does,
and this reads node boxes, edges and clusters back out of it. Each node becomes
a real shape and each edge a straight connector glued to the facing sides of
its two nodes, so a person can drag a node in Slides and its edges follow.

Anything that is not a flowchart is shown as one picture, with no clicks.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from typing import Any

from lxml import etree
from PIL import Image
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from narratr.deck import glyphs
from narratr.deck.draw import BODY_H, BODY_TOP, BODY_W, MARGIN, arrowhead, fit, panel, picture, text
from narratr.deck.layout import Ctx
from narratr.render import (
	ICON_FIELD,
	ICON_NODE,
	REMOTION,
	inline_mermaid_icons,
	mermaid_to_svg,
	run_remotion,
)
from narratr.spec import is_flowchart

SVG = "{http://www.w3.org/2000/svg}"
TRANSLATE = re.compile(r"translate\(\s*([-\d.e]+)[ ,]+([-\d.e]+)\s*\)")
# Largest scale, in inches per SVG unit. Mermaid's 20 px font at 0.025 in/px
# is 36 pt; past this a two-node diagram fills the slide with giant boxes.
MAX_SCALE = 0.025

# The connection sites python-pptx places for rectangles and diamonds.
TOP, LEFT, BOTTOM, RIGHT = 0, 1, 2, 3


@dataclass
class Node:
	id: str
	cx: float
	cy: float
	w: float
	h: float
	label: str
	kind: str  # "rect" or "diamond"
	dim: bool


@dataclass
class Edge:
	src: str
	dst: str
	label: str
	label_box: tuple[float, float, float, float] | None  # cx, cy, w, h


@dataclass
class Cluster:
	label: str
	x: float
	y: float
	w: float
	h: float


@dataclass
class Layout:
	width: float
	height: float
	nodes: dict[str, Node]
	edges: list[Edge]
	clusters: list[Cluster]


def _offset(element: Any) -> tuple[float, float]:
	"""Sum of every translate() from the element up to the root.

	Mermaid nests a subgraph's contents in their own translated group, so a
	node's own transform alone is not where it is drawn.
	"""
	x = y = 0.0
	while element is not None:
		match = TRANSLATE.search(element.get("transform") or "")
		if match:
			x += float(match.group(1))
			y += float(match.group(2))
		element = element.getparent()
	return x, y


def _text(element: Any) -> str:
	return " ".join("".join(element.itertext()).split())


def _classes(element: Any) -> list[str]:
	return (element.get("class") or "").split()


def split_edge(data_id: str, known: set[str]) -> tuple[str, str] | None:
	"""`L_<from>_<to>_<n>`, where node ids may themselves contain underscores.

	Same reading as the renderer's `edgeBeat`: try every cut and keep the one
	where both halves are nodes.
	"""
	body = re.sub(r"_\d+$", "", data_id.removeprefix("L_"))
	parts = body.split("_")
	for cut in range(1, len(parts)):
		src, dst = "_".join(parts[:cut]), "_".join(parts[cut:])
		if src in known and dst in known:
			return src, dst
	return None


def _node(g: Any, node_id: str) -> Node:
	ox, oy = _offset(g)
	shape = next((c for c in g if c.tag in {f"{SVG}rect", f"{SVG}polygon"}), None)
	kind, cx, cy, w, h = "rect", ox, oy, 0.0, 0.0
	if shape is not None and shape.tag == f"{SVG}rect":
		x, y = float(shape.get("x", 0)), float(shape.get("y", 0))
		w, h = float(shape.get("width", 0)), float(shape.get("height", 0))
		cx, cy = ox + x + w / 2, oy + y + h / 2
	elif shape is not None:
		points = [tuple(map(float, p.split(","))) for p in shape.get("points", "").split()]
		xs, ys = [p[0] for p in points], [p[1] for p in points]
		w, h = max(xs) - min(xs), max(ys) - min(ys)
		sx, sy = _offset(shape)
		cx, cy = sx + min(xs) + w / 2, sy + min(ys) + h / 2
		kind = "diamond" if len(points) == 4 else "rect"
	label = next((_text(s) for s in g.iter() if "nodeLabel" in _classes(s)), node_id)
	return Node(node_id, cx, cy, w, h, label, kind, "dim" in _classes(g))


def parse(svg: str) -> Layout:
	root = etree.fromstring(svg.encode())
	_, _, width, height = (float(v) for v in root.get("viewBox").split())
	prefix = f"{root.get('id')}-flowchart-"
	groups = list(root.iter(f"{SVG}g"))

	nodes: dict[str, Node] = {}
	for g in groups:
		gid = g.get("id") or ""
		if "node" in _classes(g) and gid.startswith(prefix):
			node_id = gid[len(prefix) :].rsplit("-", 1)[0]
			nodes[node_id] = _node(g, node_id)

	labels: dict[str, tuple[str, tuple[float, float, float, float]]] = {}
	for g in groups:
		if "label" in _classes(g) and g.get("data-id") and _text(g):
			fo = g.find(f"{SVG}foreignObject")
			ox, oy = _offset(g)
			w = float(fo.get("width", 0)) if fo is not None else 0.0
			h = float(fo.get("height", 0)) if fo is not None else 0.0
			labels[g.get("data-id")] = (_text(g), (ox + w / 2, oy + h / 2, w, h))

	edges: list[Edge] = []
	for path in root.iter(f"{SVG}path"):
		if "flowchart-link" not in _classes(path):
			continue
		data_id = path.get("data-id") or ""
		ends = split_edge(data_id, set(nodes))
		if ends is not None:
			caption, box = labels.get(data_id, ("", None))
			edges.append(Edge(ends[0], ends[1], caption, box))

	clusters: list[Cluster] = []
	for g in groups:
		rect = g.find(f"{SVG}rect")
		if "cluster" not in _classes(g) or rect is None:
			continue
		ox, oy = _offset(g)
		label = next((_text(s) for s in g.iter() if "cluster-label" in _classes(s)), "")
		clusters.append(
			Cluster(
				label,
				ox + float(rect.get("x", 0)),
				oy + float(rect.get("y", 0)),
				float(rect.get("width", 0)),
				float(rect.get("height", 0)),
			)
		)
	return Layout(width, height, nodes, edges, clusters)


def _icons(source: str) -> dict[str, str]:
	found: dict[str, str] = {}
	for match in ICON_NODE.finditer(source):
		icon = ICON_FIELD.search(match.group("body"))
		if icon:
			found[match.group("id")] = icon.group("name")
	return found


def _font_px(source: str) -> float:
	match = re.search(r"fontSize:\s*(\d+)", source)
	return float(match.group(1)) if match else 20.0


def _sides(a: Node, b: Node) -> tuple[int, int]:
	dx, dy = b.cx - a.cx, b.cy - a.cy
	if abs(dx) >= abs(dy):
		return (RIGHT, LEFT) if dx >= 0 else (LEFT, RIGHT)
	return (BOTTOM, TOP) if dy >= 0 else (TOP, BOTTOM)


def _picture(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	"""Not a flowchart: mmdc's own PNG, themed to match the deck, shown whole."""
	theme = ctx.theme
	config = ctx.tmp / "deck-mermaid.json"
	config.write_text(
		json.dumps(
			{
				"theme": "base",
				"securityLevel": "strict",
				"themeVariables": {
					"darkMode": not theme.is_light,
					"background": theme.background,
					"mainBkg": theme.panel,
					"primaryColor": theme.panel,
					"primaryTextColor": theme.text,
					"primaryBorderColor": theme.line,
					"lineColor": theme.line,
					"textColor": theme.text,
					"fontFamily": theme.body_font,
					"fontSize": "20px",
				},
			}
		)
	)
	mmd = ctx.tmp / f"{scene['id']}.mmd"
	png = ctx.tmp / f"{scene['id']}.png"
	mmd.write_text(inline_mermaid_icons(scene["mermaid"]))
	run_remotion(
		[
			*("npx", "--no-install", "mmdc", "-i", str(mmd), "-o", str(png)),
			*("-b", "transparent", "-c", str(config), "-s", "3"),
		],
		REMOTION,
		"mermaid",
	)
	with Image.open(png) as img:
		iw, ih = img.size
	picture(slide.shapes, png.read_bytes(), *fit(iw, ih, MARGIN, BODY_TOP, BODY_W, BODY_H))
	return []


def draw(slide: Any, scene: dict[str, Any], ctx: Ctx) -> list[list[int]]:
	source = scene["mermaid"]
	if not is_flowchart(source):
		return _picture(slide, scene, ctx)
	theme = ctx.theme
	layout = parse(mermaid_to_svg(source, ctx.tmp / f"{scene['id']}.svg"))
	scale = min(BODY_W / layout.width, BODY_H / layout.height, MAX_SCALE)
	ox = MARGIN + (BODY_W - layout.width * scale) / 2
	oy = BODY_TOP + (BODY_H - layout.height * scale) / 2
	font = max(12.0, min(32.0, _font_px(source) * scale * 72))
	small = max(12.0, font * 0.8)
	icon_for = _icons(source)

	def at(x: float, y: float) -> tuple[float, float]:
		return ox + x * scale, oy + y * scale

	# Clusters first, so they sit behind everything. They never animate.
	for cluster in layout.clusters:
		x, y = at(cluster.x, cluster.y)
		w, h = cluster.w * scale, cluster.h * scale
		panel(slide.shapes, x, y, w, h, fill=theme.background, border=theme.line)
		if cluster.label:
			text(
				slide.shapes, x + 0.1, y + 0.05, w - 0.2, 0.4, cluster.label,
				size=small, colour=theme.muted, font=theme.body_font,
			)  # fmt: skip

	steps = scene.get("revealSteps", [])
	beat = {node: i for i, step in enumerate(steps) for node in step}
	groups: list[list[int]] = [[] for _ in steps]
	shapes: dict[str, Any] = {}
	for index, node in enumerate(layout.nodes.values()):
		x, y = at(node.cx - node.w / 2, node.cy - node.h / 2)
		w, h = node.w * scale, node.h * scale
		border = theme.line if node.dim else theme.accent(index)
		kind = MSO_SHAPE.DIAMOND if node.kind == "diamond" else MSO_SHAPE.ROUNDED_RECTANGLE
		box = panel(slide.shapes, x, y, w, h, fill=theme.panel, border=border, shape=kind)
		box.name = f"node:{node.id}"
		shapes[node.id] = box
		members = [box.shape_id]
		inset = 0.0
		icon = icon_for.get(node.id)
		if icon:
			size = min(h * 0.5, font / 72 * 1.2)
			glyph = glyphs.png(icon, border)
			members.append(
				picture(slide.shapes, glyph, x + 0.15, y + (h - size) / 2, size, size).shape_id
			)
			inset = size + 0.15
		label = text(
			slide.shapes, x + inset + 0.08, y, w - inset - 0.16, h, node.label,
			size=font, colour=theme.muted if node.dim else theme.text, font=theme.body_font,
			align=PP_ALIGN.CENTER,
		)  # fmt: skip
		members.append(label.shape_id)
		if node.id in beat:
			groups[beat[node.id]] += members

	for edge in layout.edges:
		start, end = _sides(layout.nodes[edge.src], layout.nodes[edge.dst])
		line = slide.shapes.add_connector(
			MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(0), Inches(1), Inches(1)
		)
		line.begin_connect(shapes[edge.src], start)
		line.end_connect(shapes[edge.dst], end)
		arrowhead(line, theme.line)
		line.name = f"edge:{edge.src}->{edge.dst}"
		members = [line.shape_id]
		if edge.label and edge.label_box:
			cx, cy, w, h = edge.label_box
			x, y = at(cx - w / 2, cy - h / 2)
			caption = text(
				slide.shapes, x, y, w * scale, h * scale, edge.label,
				size=small, colour=theme.muted, font=theme.body_font, align=PP_ALIGN.CENTER,
			)  # fmt: skip
			members.append(caption.shape_id)
		# An edge arrives with the later of its two nodes; between two nodes
		# that never reveal, it is there from the start.
		when = max(beat.get(edge.src, -1), beat.get(edge.dst, -1))
		if when >= 0:
			groups[when] += members
	return groups
