"""Drawing primitives shared by every layout. Positions are in inches."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from narratr.deck.theme import rgb

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.6
BODY_TOP = 1.6
BODY_W = SLIDE_W - 2 * MARGIN
BODY_H = SLIDE_H - BODY_TOP - 0.5


def textbox(shapes: Any, x: float, y: float, w: float, h: float, anchor: MSO_ANCHOR) -> Any:
	"""An empty, fixed-size box: no autofit, so editing it never reflows the slide."""
	box = shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
	frame = box.text_frame
	frame.word_wrap = True
	frame.auto_size = None
	frame.vertical_anchor = anchor
	for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
		setattr(frame, side, Emu(0))
	return box


def text(
	shapes: Any,
	x: float,
	y: float,
	w: float,
	h: float,
	content: str,
	*,
	size: float,
	colour: str,
	font: str,
	bold: bool = False,
	align: PP_ALIGN = PP_ALIGN.LEFT,
	anchor: MSO_ANCHOR = MSO_ANCHOR.MIDDLE,
	strike: bool = False,
) -> Any:
	box = textbox(shapes, x, y, w, h, anchor)
	para = box.text_frame.paragraphs[0]
	para.alignment = align
	run = para.add_run()
	run.text = content
	style(run, size=size, colour=colour, font=font, bold=bold, strike=strike)
	return box


def style(
	run: Any, *, size: float, colour: str, font: str, bold: bool = False, strike: bool = False
) -> None:
	run.font.size = Pt(size)
	run.font.bold = bold
	run.font.name = font
	run.font.color.rgb = rgb(colour)
	if strike:
		run.font._rPr.set("strike", "sngStrike")


def panel(
	shapes: Any,
	x: float,
	y: float,
	w: float,
	h: float,
	*,
	fill: str,
	border: str | None = None,
	shape: MSO_SHAPE = MSO_SHAPE.ROUNDED_RECTANGLE,
) -> Any:
	box = shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
	box.fill.solid()
	box.fill.fore_color.rgb = rgb(fill)
	if border:
		box.line.color.rgb = rgb(border)
		box.line.width = Pt(1.5)
	else:
		box.line.fill.background()
	box.shadow.inherit = False
	if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
		box.adjustments[0] = 0.08
	return box


def picture(shapes: Any, data: bytes, x: float, y: float, w: float, h: float) -> Any:
	return shapes.add_picture(BytesIO(data), Inches(x), Inches(y), Inches(w), Inches(h))


def arrow(shapes: Any, x1: float, y1: float, x2: float, y2: float, *, colour: str) -> Any:
	line = shapes.add_connector(
		MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
	)
	arrowhead(line, colour)
	return line


def arrowhead(connector: Any, colour: str) -> None:
	connector.line.color.rgb = rgb(colour)
	connector.line.width = Pt(2)
	ln = connector.line._get_or_add_ln()
	ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def fit(
	img_w: float, img_h: float, x: float, y: float, w: float, h: float
) -> tuple[float, float, float, float]:
	"""Largest box of the image's aspect inside (x, y, w, h), centred."""
	scale = min(w / img_w, h / img_h)
	fw, fh = img_w * scale, img_h * scale
	return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh
