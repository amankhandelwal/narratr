"""scenes.json -> deck.pptx + script.md."""

from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from narratr.deck import animate, diagram, script
from narratr.deck.draw import MARGIN, SLIDE_H, SLIDE_W
from narratr.deck.layout import LAYOUTS, Ctx
from narratr.deck.theme import Theme, resolve, rgb
from narratr.errors import PipelineError

# Registered here rather than in layout.py, which diagram.py imports from.
LAYOUTS["diagram"] = diagram.draw

# python-pptx's default template: 0 is "Title Slide", 5 is "Title Only".
SECTION_LAYOUT = 0
CONTENT_LAYOUT = 5


class DeckError(PipelineError):
	"""The deck could not be built."""


def _theme_fonts(prs: Any, theme: Theme) -> None:
	"""Point the theme's heading and body fonts at ours, so text a person adds
	by hand later, in Slides or PowerPoint, matches what was generated."""
	part = prs.slide_master.part.part_related_by(RT.THEME)
	root = etree.fromstring(part.blob)
	for which, font in (("a:majorFont", theme.heading_font), ("a:minorFont", theme.body_font)):
		latin = root.find(f".//{qn(which)}/{qn('a:latin')}")
		if latin is not None:
			latin.set("typeface", font)
	part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _master(prs: Any, theme: Theme) -> None:
	prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
	fill = prs.slide_master.background.fill
	fill.solid()
	fill.fore_color.rgb = rgb(theme.background)
	_theme_fonts(prs, theme)


def _heading(slide: Any, scene: dict[str, Any], theme: Theme, section: bool) -> None:
	title = slide.shapes.title
	heading = scene.get("heading")
	# An unused placeholder shows "Click to add text" to whoever edits the deck.
	for placeholder in list(slide.placeholders):
		if placeholder.placeholder_format.idx != 0 or not heading:
			placeholder._element.getparent().remove(placeholder._element)
	if not heading:
		return
	top, height = (2.5, 2.5) if section else (0.4, 1.0)
	title.left, title.top = Inches(MARGIN), Inches(top)
	title.width, title.height = Inches(SLIDE_W - 2 * MARGIN), Inches(height)
	frame = title.text_frame
	frame.text = heading
	para = frame.paragraphs[0]
	para.alignment = PP_ALIGN.CENTER if section else PP_ALIGN.LEFT
	run = para.runs[0]
	run.font.size = Pt(60 if section else 40)
	run.font.bold = True
	run.font.name = theme.heading_font
	run.font.color.rgb = rgb(theme.text)


def is_section(scene: dict[str, Any]) -> bool:
	"""A heading and nothing else: a divider between parts of the talk."""
	return scene["type"] == "prose" and not scene.get("bullets") and not scene.get("image")


def build(spec: dict[str, Any], base: Path, out_dir: Path) -> tuple[Path, Path]:
	"""Write deck.pptx and script.md into out_dir. The spec must already be valid."""
	theme = resolve(spec)
	prs = Presentation()
	_master(prs, theme)
	with tempfile.TemporaryDirectory() as tmp:
		ctx = Ctx(theme=theme, base=base, tmp=Path(tmp))
		for scene in spec["scenes"]:
			draw = LAYOUTS.get(scene["type"])
			if draw is None:
				raise DeckError(f"scene '{scene['id']}': no deck layout for '{scene['type']}'")
			section = is_section(scene)
			layout = prs.slide_layouts[SECTION_LAYOUT if section else CONTENT_LAYOUT]
			slide = prs.slides.add_slide(layout)
			_heading(slide, scene, theme, section)
			groups = draw(slide, scene, ctx)
			# The invariant every deck test leans on, checked on every build
			# rather than only on the bundled examples.
			expected = len(script.reveals(scene))
			if len(groups) != expected:
				raise DeckError(
					f"scene '{scene['id']}': drew {len(groups)} reveals for {expected} cues"
				)
			animate.apply(slide, groups)
			slide.notes_slide.notes_text_frame.text = script.marked(scene)
	out_dir.mkdir(parents=True, exist_ok=True)
	deck, doc = out_dir / "deck.pptx", out_dir / "script.md"
	prs.save(str(deck))
	doc.write_text(script.document(spec))
	return deck, doc
